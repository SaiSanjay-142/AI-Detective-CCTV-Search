"""Create a transparent evaluation-readiness report from the MAGNAS SQLite evidence store."""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "ai_detective.db"
OUT = ROOT / "output" / "evaluation" / "forensic_video_evaluation_report.docx"


def query(cursor: sqlite3.Cursor, sql: str):
    return cursor.execute(sql).fetchall()


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}"); node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "4"); node.set(qn("w:color"), "D9D9D9"); borders.append(node)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = header; shade(cell, "17365D"); border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(8.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths: cell.width = Inches(widths[i])
    for row_num, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value); border(cells[i]); cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_num % 2: shade(cells[i], "EEF3F8")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs: run.font.size = Pt(8.5)
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(5)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs: run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(4)


def pct(n: int, d: int) -> str:
    return "N/A" if not d else f"{100 * n / d:.1f}%"


def main() -> None:
    conn = sqlite3.connect(DB); c = conn.cursor()
    statuses = dict(query(c, "select status,count(*) from processing_jobs group by status"))
    videos = c.execute("select count(*) from videos").fetchone()[0]
    frames = c.execute("select count(*) from frames").fetchone()[0]
    detections, det_avg, det_min, det_max = c.execute("select count(*),avg(confidence),min(confidence),max(confidence) from detections").fetchone()
    objects, obj_avg, obj_min, obj_max = c.execute("select count(*),avg(confidence),min(confidence),max(confidence) from object_observations").fetchone()
    tracks, track_avg, track_min, track_max = c.execute("select count(*),avg(detection_count),min(detection_count),max(detection_count) from tracks").fetchone()
    attrs, gender, upper_type, upper_colour, lower_type, lower_colour = c.execute("select count(*),sum(gender != 'unknown'),sum(upper_type != 'unknown'),sum(upper_color != 'unknown'),sum(lower_type != 'unknown'),sum(lower_color != 'unknown') from person_attributes").fetchone()
    actions, action_avg, action_min, action_max = c.execute("select count(*),avg(confidence),min(confidence),max(confidence) from actions").fetchone()
    audits, parse_errors = c.execute("select count(*),sum(parse_error is not null) from attribute_audits").fetchone()
    searches = c.execute("select count(*) from search_queries").fetchone()[0]
    top_objects = query(c, "select label,count(*),avg(confidence) from object_observations group by label order by count(*) desc limit 10")
    conn.close()

    doc = Document(); section = doc.sections[0]; section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65); section.left_margin = Inches(0.7); section.right_margin = Inches(0.7)
    styles = doc.styles; styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CCTV Forensic Search Evaluation Report"); run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Database evidence snapshot and 20 epoch evaluation readiness").italic = True
    subtitle.runs[0].font.size = Pt(10)
    doc.add_paragraph(f"Generated {datetime.now().strftime('%d %B %Y %H:%M')} from ai_detective.db. The database contains {videos} uploaded video records and {frames} indexed sampled frames.")

    add_heading(doc, "Conclusion")
    doc.add_paragraph("The current evidence store supports operational coverage measurement, but it does not contain human-labelled bounding boxes, person attributes, actions, or question-answer ground truth. As a result, detection mAP, detection precision and recall, attribute accuracy, action accuracy, and Gemini answer accuracy cannot be calculated validly. A 20 epoch YOLO training and validation run was not executed because no labelled train, validation, and test split is present. The values below are measured system-output statistics, not accuracy claims.")

    add_heading(doc, "Evaluation scope and data readiness")
    add_table(doc, ["Area", "Available evidence", "Accuracy metric status"], [
        ["YOLO object and person detection", f"{detections:,} person detections; {objects:,} non-person observations", "mAP, precision and recall: not measurable without annotated boxes"],
        ["ByteTrack person tracking", f"{tracks:,} tracks; {track_avg:.2f} mean sampled detections per track", "MOTA, IDF1 and HOTA: not measurable without identity annotations"],
        ["Qwen PP Human attributes", f"{attrs:,} stored person-attribute records; {audits:,} model audits", "Per-field accuracy and F1: not measurable without attribute labels"],
        ["VideoMAE actions", f"{actions:,} action segments stored", "Action accuracy: not measurable without action labels"],
        ["Gemini retrieval answers", f"{searches:,} saved search queries", "Answer correctness: not measurable because expected answers and response audit are not stored"],
        ["20 epoch training", "No YOLO dataset YAML or label files found", "Not run - no supervised training split available"],
    ], [1.45, 2.55, 2.7])

    add_heading(doc, "Observed processing output")
    add_table(doc, ["Measure", "Value", "Interpretation"], [
        ["Processing jobs", f"{sum(statuses.values())} total; {statuses.get('completed',0)} completed; {statuses.get('completed_with_warnings',0)} with warnings; {statuses.get('processing',0)} still processing", "Job state snapshot; not a model-quality score"],
        ["Person detection confidence", f"mean {det_avg:.3f}; range {det_min:.3f} to {det_max:.3f}", "Model confidence, not precision"],
        ["Object detection confidence", f"mean {obj_avg:.3f}; range {obj_min:.3f} to {obj_max:.3f}", "Model confidence, not precision"],
        ["Track observations", f"mean {track_avg:.2f}; range {track_min} to {track_max}", "Evidence density per ByteTrack ID"],
        ["Action confidence", f"mean {action_avg:.3f}; range {action_min:.3f} to {action_max:.3f}", "Model confidence; requires labelled validation"],
        ["Attribute parse failures", f"{parse_errors or 0} of {audits} audits", "Format/availability reliability, not attribute correctness"],
    ], [1.75, 1.8, 3.15])

    add_heading(doc, "Attribute retrieval coverage")
    add_table(doc, ["Indexed attribute", "Known values", "Coverage among stored attribute records", "Accuracy"], [
        ["Gender presentation", str(gender or 0), pct(gender or 0, attrs), "Not measurable without labels"],
        ["Upper garment type", str(upper_type or 0), pct(upper_type or 0, attrs), "Not measurable without labels"],
        ["Upper garment colour", str(upper_colour or 0), pct(upper_colour or 0, attrs), "Not measurable without labels"],
        ["Lower garment type", str(lower_type or 0), pct(lower_type or 0, attrs), "Not measurable without labels"],
        ["Lower garment colour", str(lower_colour or 0), pct(lower_colour or 0, attrs), "Not measurable without labels"],
        ["Track to attribute coverage", f"{attrs} of {tracks}", pct(attrs, tracks), "Completeness only; a track may lack a suitable crop"],
    ], [1.65, 1.15, 1.65, 2.25])

    add_heading(doc, "Most frequent indexed object observations")
    add_table(doc, ["Object class", "Frame observations", "Mean detector confidence"], [[label, str(count), f"{confidence:.3f}"] for label, count, confidence in top_objects], [2.35, 1.75, 2.6])

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Required 20 epoch evaluation protocol")
    doc.add_paragraph("To produce real accuracy values, create a non-overlapping train, validation, and test split from the CCTV footage. Label each sampled frame with person and object bounding boxes in YOLO format. For every representative person crop, add reviewed labels for the permitted attribute taxonomy. Create a separate question-answer set where every question links to the expected tracks, frames, and answer.")
    add_table(doc, ["Component", "20 epoch procedure", "Metrics to report"], [
        ["YOLO detection", "Fine-tune on training split for 20 epochs; select the best validation checkpoint; run once on held-out test set.", "mAP50, mAP50-95, precision, recall, per-class AP and confusion matrix"],
        ["Tracking", "Run fixed detector checkpoint and ByteTrack configuration on annotated test videos.", "HOTA, IDF1, MOTA, ID switches and mostly tracked rate"],
        ["Attributes", "Freeze the test set before threshold selection; compare each visible attribute with reviewer labels.", "Per-field accuracy, macro F1, precision, recall, coverage and unknown rate"],
        ["Actions", "Evaluate temporal segments only against action-labelled test clips.", "Top-1 accuracy, macro F1, segment IoU and per-action recall"],
        ["Gemini retrieval", "Use a held-out set of evidence-grounded questions; manually adjudicate answer and selected evidence.", "Answer exact match or rubric score, evidence precision, evidence recall and unsupported-claim rate"],
    ], [1.2, 3.25, 2.25])

    add_heading(doc, "Acceptance criteria")
    doc.add_paragraph("Report all confidence intervals or test-set sample counts beside each accuracy measure. Do not use detector confidence as accuracy. Keep train, validation, and test videos separate to prevent near-duplicate frames from inflating results. For gender presentation, clothing colour, and other sensitive visual attributes, retain an unknown class and report abstention/coverage separately from accuracy.")

    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
