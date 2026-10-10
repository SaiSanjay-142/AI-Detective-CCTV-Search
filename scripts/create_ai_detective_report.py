from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_ORIENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "AI_Detective_Project_Report.docx"
OUT.parent.mkdir(exist_ok=True)

NAVY = "17365D"; BLUE = "2F75B5"; PALE = "EAF1F8"; LIGHT = "F4F7FA"; GRID = "D9D9D9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), GRID); b.append(e)
    tblPr.append(b)

def set_cell(cell, text, bold=False, color=None, size=9.5, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    if align is not None: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for side in ('top','left','bottom','right'):
        x = OxmlElement(f'w:{side}'); x.set(qn('w:w'),'100'); x.set(qn('w:type'),'dxa'); mar.append(x)
    tcPr.append(mar)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); set_cell(c,h,True,"FFFFFF",9,WD_ALIGN_PARAGRAPH.CENTER)
        if widths: c.width=Inches(widths[i])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            if ri%2==1: shade(cells[i],LIGHT)
            set_cell(cells[i],str(val),False,None,9.3, WD_ALIGN_PARAGRAPH.CENTER if i==0 and len(str(val))<22 else None)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.font.color.rgb=RGBColor(0,0,0); return p

def para(doc, text, bold_lead=None):
    p=doc.add_paragraph(style="Normal"); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.15
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else:p.add_run(text)
    return p

def bullets(doc, items):
    for x in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3); p.add_run(x)

def page_number(section):
    footer=section.footer; p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("AI Detective Project Report | ")
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
page_number(sec)

styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); styles['Normal'].font.size=Pt(10.5)
for name,size in [('Title',24),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    s=styles[name]; s.font.name='Aptos Display' if name!='Normal' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'),s.font.name); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)

# Cover
for _ in range(5): doc.add_paragraph()
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('AI Detective')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Vision Language CCTV Investigation and Search System'); r.bold=True; r.font.size=Pt(17)
for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('PROJECT REPORT').bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Phase I Technical Documentation')
for _ in range(6): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Prepared for project review').italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('October 2026')
doc.add_page_break()

heading(doc,'Abstract',1)
para(doc,'AI Detective is a web based CCTV investigation system that converts uploaded surveillance footage into searchable visual evidence. The system samples video frames, detects people using YOLOv8, maintains identities with ByteTrack, extracts visible attributes through Qwen3 VL, and is designed to recognize actions through VideoMAE. Structured metadata is stored in SQLite and queried through an explainable natural language parser. Instead of returning only a frame, the intended result contains a track identifier, timestamp interval, evidence frame, bounding box, and the reasons for the match.')
para(doc,'The implemented application provides local single user upload, real video metadata extraction, sampled frame storage, YOLOv8 person detection, ByteTrack tracking, cropped evidence storage, SQLite indexing, deterministic query parsing, temporal evidence retrieval, and original video seeking. Qwen3 VL is integrated as an attribute extraction component but depends on sufficient CUDA capable hardware and model availability. VideoMAE requires a compatible action classification checkpoint and its label mapping before action results can be reported. This report therefore separates implemented capabilities from pending model dependent capabilities and does not claim unmeasured accuracy.')
heading(doc,'Keywords',2); para(doc,'CCTV investigation, computer vision, YOLOv8, ByteTrack, Qwen3 VL, VideoMAE, FastAPI, React, SQLite, temporal evidence retrieval.')

heading(doc,'Table of Contents',1)
for item in ['1 Introduction','2 Objectives and Scope','3 Literature and Technology Background','4 System Design','5 Methodology and Processing Pipeline','6 Implementation Details','7 Natural Language Search and Evidence Retrieval','8 Metrics and Evaluation Plan','9 Current Results and Limitations','10 Conclusion and Future Work','References','Appendix A API Summary']:
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); p.add_run(item)
doc.add_page_break()

heading(doc,'1 Introduction',1)
heading(doc,'1.1 Problem Statement',2)
para(doc,'Investigators often need to inspect long CCTV recordings to locate a person, object, or action. Manual review is slow and makes it difficult to compare people across time. AI Detective addresses this by converting selected visual evidence into structured, searchable metadata while retaining the original video for human verification.')
heading(doc,'1.2 Proposed Solution',2)
para(doc,'The project combines detection, tracking, visual attribute extraction, action recognition, structured storage, and natural language retrieval. A user uploads a video, monitors actual processing progress, enters a query such as "find a person with a backpack", and receives matching tracks with timestamps and evidence frames. The original video remains visible so the user can verify the model output.')
heading(doc,'1.3 Objectives',2)
bullets(doc,['Accept common CCTV video formats and record the video FPS, duration, width, and height.','Detect people using a real YOLOv8 model and record bounding boxes with confidence scores.','Assign stable track identifiers using ByteTrack.','Extract controlled visual attributes from person and context crops using Qwen3 VL when the model runtime is available.','Support temporal action segments through VideoMAE when a compatible fine tuned checkpoint is supplied.','Provide deterministic natural language query parsing and explainable SQLite filtering.','Show evidence timestamps, actual frames, bounding boxes, and original video seeking.'])

heading(doc,'2 Objectives and Scope',1)
table(doc,['Area','In Scope','Current Status'],[
['Video ingestion','MP4 AVI MOV MKV validation, safe storage, metadata extraction','Implemented'],
['Detection and tracking','YOLOv8 person detection and ByteTrack track IDs','Implemented subject to model runtime'],
['Attribute analysis','Qwen3 VL structured JSON for clothing, accessories and held objects','Integrated; hardware dependent'],
['Action analysis','VideoMAE temporal action windows and intervals','Requires compatible checkpoint'],
['Search','Rule based query parsing and SQLite set intersection','Implemented'],
['Evidence','Actual stored frame, overlay box, temporal seek in original video','Implemented'],
['Multi user deployment','JWT endpoints retained, local UI uses one local workspace owner','Not primary scope']],[1.3,3.4,1.5])
heading(doc,'2.1 Scope Boundary',2)
para(doc,'AI Detective is an investigative assistance system. It reports model detected visual evidence and must not be used to claim that a person committed an offence. Every result must be verified against the original footage. The project does not fabricate detections, attributes, actions, confidence values, or benchmark scores.')

heading(doc,'3 Literature and Technology Background',1)
heading(doc,'3.1 YOLOv8 for Detection',2)
para(doc,'YOLOv8 is a one stage object detector. A frame passes through a feature extraction backbone, a multi scale feature fusion neck, and a detection head that predicts object class, confidence, and bounding box coordinates. In this project the detector is configured to prioritize the person class. A detection record stores video ID, frame ID, track ID, bounding box, label, and confidence.')
heading(doc,'3.2 ByteTrack for Multi Object Tracking',2)
para(doc,'ByteTrack links detections between consecutive frames. It predicts existing track positions and matches them with new detections using motion and spatial overlap. Its useful property is two stage association: high confidence detections establish strong matches, while lower confidence detections can recover a temporarily occluded or blurred person. This reduces fragmented trajectories without adding a separate re identification model.')
heading(doc,'3.3 Qwen3 VL for Attribute Extraction',2)
para(doc,'Qwen3 VL is a vision language model. The project provides a narrow person crop and a context crop. Its visual encoder converts image content to visual tokens, which are combined with the CCTV instruction prompt in the language model. The prompt requests only structured JSON and requires unknown when the image cannot support a reliable value. Pydantic validation prevents uncontrolled model output from entering the database.')
heading(doc,'3.4 VideoMAE for Action Recognition',2)
para(doc,'VideoMAE is a video transformer that models motion across a sequence rather than a single image. Frames are divided into patches and temporal tubelets, transformed into tokens, and processed by attention layers. A fine tuned classification head produces action probabilities for a temporal window. For this project, overlapping 16 frame windows with stride 8 are planned, then adjacent windows with the same action are merged into one start time and end time interval.')

heading(doc,'4 System Design',1)
heading(doc,'4.1 Architecture Overview',2)
para(doc,'The frontend is built with React, TypeScript, and Vite. It presents upload, progress, a persistent original video player, query input, temporal results, and evidence frames. The backend is built with FastAPI and Python. SQLite is used for structured evidence metadata. OpenCV reads video frames and generates cropped images. Ultralytics provides YOLOv8 and ByteTrack integration. Hugging Face Transformers provides the Qwen3 VL and VideoMAE model interfaces.')
table(doc,['Layer','Technology','Responsibility'],[
['Presentation','React TypeScript Vite','Upload video, display job progress, query evidence, seek original video'],
['API','FastAPI Uvicorn','Upload, processing, search, evidence and local user endpoints'],
['Processing','OpenCV Python','Read video, sample frames, save crops and overlays'],
['Detection and tracking','YOLOv8m or YOLOv8x, ByteTrack','Find people and maintain track IDs'],
['Vision language','Qwen3 VL 2B Instruct','Extract controlled person attributes'],
['Action recognition','VideoMAE Large','Classify temporal action windows when checkpoint is supplied'],
['Persistence','SQLite SQLAlchemy','Store videos, frames, detections, tracks, attributes, actions and searches']],[1.25,2.25,2.7])
heading(doc,'4.2 Data Flow',2)
bullets(doc,['Upload video -> validate format -> store original video and metadata.','Create background processing job -> sample frames using real video FPS.','YOLOv8 detection -> ByteTrack association -> store detections and tracks.','Create person and context crops -> Qwen3 VL attribute extraction where available.','Create action windows -> VideoMAE action classification where available.','Store structured metadata -> parse natural language query -> filter matching tracks and event intervals.','Return evidence frame and interval -> seek original video to the exact start time.'])

heading(doc,'5 Methodology and Processing Pipeline',1)
heading(doc,'5.1 Video Ingestion and Frame Sampling',2)
para(doc,'OpenCV opens the uploaded source and reads its actual frames per second, frame count, dimensions, and duration. The application derives sampling cadence from FRAME_SAMPLE_FPS, which defaults to 3 FPS. Timestamps are calculated using the actual source FPS for constant frame rate material. Each stored frame preserves its video ID, original frame number, timestamp, and image path.')
heading(doc,'5.2 Detection Tracking and Cropping',2)
para(doc,'Each sampled frame is passed to YOLOv8 tracking mode with ByteTrack enabled. The system stores each person bounding box and associates it with a track ID. The first suitable observation for a new track creates a narrow crop and a context crop with 150 pixels of padding where frame boundaries allow. The narrow crop supports clothing analysis; the context crop supports bags, bottles, phones, and nearby objects.')
heading(doc,'5.3 Attribute Extraction',2)
para(doc,'The attribute extractor loads Qwen3 VL lazily so the web server can start without loading the multimodal model. It sends a constrained prompt and the two crops, extracts JSON from the generated response, and validates a fixed schema. Allowed values include controlled gender, clothing type, colour, accessories, headwear, eyewear, and held objects. A malformed response is isolated and reported as a warning rather than stopping the entire processing job.')
heading(doc,'5.4 Action Recognition',2)
para(doc,'The action recognizer is intentionally guarded. It checks ACTION_MODEL_PATH and does not create action labels if a compatible VideoMAE checkpoint is missing. A correct action checkpoint must contain a classification head and label mapping for the project action vocabulary. Once supplied, the implementation should load each temporal window, normalize it with the checkpoint processor, infer label probabilities, apply confidence thresholding and smoothing, merge neighbouring action windows, and store action intervals.')

heading(doc,'6 Implementation Details',1)
table(doc,['Module','File','Function'],[
['Application setup','backend/app/main.py','FastAPI startup, CORS, routes and database initialization'],
['Configuration','backend/app/config.py','Environment variables, model names, device, paths and frame sampling rate'],
['Database models','backend/app/models.py','SQLite tables and relationships'],
['Video API','backend/app/api/videos.py','Upload, list, stream, process and evidence frame endpoints'],
['Processing API','backend/app/api/processing.py','Actual background job status and progress endpoint'],
['Search API','backend/app/api/search.py','Natural language search and evidence endpoints'],
['Pipeline','backend/app/services/video_processor.py','Frame processing, detection, storage, crops, Qwen stage and progress updates'],
['Detector','backend/app/services/detector.py','YOLOv8 and ByteTrack call'],
['Attributes','backend/app/services/attribute_extractor.py','Qwen prompt, model loading, JSON extraction and validation'],
['Actions','backend/app/services/action_recognizer.py','VideoMAE checkpoint guard and future inference interface'],
['Query parser','backend/app/services/query_parser.py','Synonyms, colours, actions and structured criteria'],
['Retrieval','backend/app/services/retrieval.py','Metadata filtering, score explanation and evidence selection'],
['Frontend','frontend/src/main.tsx','Upload, progress, video player, query and evidence interaction']],[1.2,2.6,2.4])
heading(doc,'6.1 Database Design',2)
table(doc,['Table','Purpose','Key Fields'],[
['videos','Original footage metadata','owner ID, filename, stored path, FPS, duration, dimensions'],
['processing_jobs','Actual pipeline state','status, stage, progress, detail, error'],
['frames','Sampled source frames','video ID, frame number, timestamp, image path'],
['detections','YOLO person evidence','frame ID, video ID, track ID, box, confidence'],
['tracks','Track level summary','video ID, track ID, first and last timestamp'],
['person_attributes','Validated Qwen output','gender, clothing, accessories, objects, confidence'],
['actions','VideoMAE intervals','track ID, action, start, end, confidence'],
['search_queries','Search traceability','query text and parsed JSON']],[1.3,2.3,2.6])

heading(doc,'7 Natural Language Search and Evidence Retrieval',1)
heading(doc,'7.1 Query Parsing',2)
para(doc,'The parser works without an external language API for the supported vocabulary. It normalizes variants such as man to male, woman to female, rucksack to backpack, walking to walk, running to run, and dark green to dark_green. The parsed query is stored for traceability.')
para(doc,'Example query: Find a man in a dark green hoodie with a backpack walking.')
table(doc,['Field','Parsed Criterion'],[['gender','male'],['upper clothing type','hoodie'],['upper clothing colour','dark_green'],['accessory','backpack'],['action','walk']],[2.4,3.8])
heading(doc,'7.2 Explainable Retrieval',2)
para(doc,'The retrieval engine applies each supplied criterion as a hard filter. Conceptually, the result is the intersection of matching sets: male intersection dark green hoodie intersection backpack intersection walking. After filtering, the system adds a transparent score based on matched criteria and track continuity. It does not use opaque similarity ranking for the main retrieval path.')
heading(doc,'7.3 Temporal Evidence',2)
para(doc,'For a track based attribute result, the evidence interval is the observed first and last timestamp for that track. For an action query, each matching VideoMAE action segment is returned independently. A detection within the interval closest to its midpoint becomes the evidence frame. The frontend displays start time, end time, duration and the actual evidence timestamp. Selecting a result seeks the protected original video to the event start and pauses at the event end.')

heading(doc,'8 Metrics and Evaluation Plan',1)
para(doc,'The project must be evaluated using labelled data. The present implementation does not contain a labelled CCTV test set or measured model benchmark, so numerical results must not be claimed until an evaluation dataset and ground truth annotations are available. The following metrics define the correct evaluation plan.')
heading(doc,'8.1 Detection and Tracking Metrics',2)
table(doc,['Metric','Meaning','Required Ground Truth'],[
['Precision','Fraction of detected persons that are true persons','Person bounding boxes'],
['Recall','Fraction of true persons that were detected','Person bounding boxes'],
['mAP at IoU threshold','Detection quality over precision recall curve','Boxes and class labels'],
['MOTA','Tracking accuracy with missed, false and ID switch errors','Track IDs across frames'],
['IDF1','Identity consistency across the trajectory','Track IDs across frames'],
['ID switches','Number of times a person receives a different identity','Track IDs across frames']],[1.4,2.3,2.5])
heading(doc,'8.2 Attribute and Action Metrics',2)
table(doc,['Metric','Use','Interpretation'],[
['Attribute precision recall F1','Gender, clothing, accessories and held objects','Measures correctness of structured Qwen attributes'],
['Action precision recall F1','Walk, run, pick and other configured labels','Measures predicted action class correctness'],
['Temporal IoU','Action start and end intervals','Overlap between predicted and labelled action time'],
['mAP for actions','Action detection across confidence thresholds','Ranks multiple temporal action events'],
['Calibration error','Confidence score reliability','Checks whether a 0.8 confidence behaves like 80 percent correctness']],[1.5,2.2,2.5])
heading(doc,'8.3 Retrieval and System Metrics',2)
table(doc,['Metric','Formula or Measurement','Purpose'],[
['Precision at K','relevant results in top K divided by K','Quality of first investigator results'],
['Recall at K','relevant results in top K divided by all relevant results','Coverage of known evidence'],
['Retrieval mAP','mean average precision across queries','Overall search ranking quality'],
['Query latency','request time from search submit to result response','Investigation responsiveness'],
['Processing throughput','video seconds processed per wall clock second','Pipeline speed'],
['P50 and P95 latency','median and 95th percentile processing or query time','Stability under repeated use']],[1.6,2.5,2.1])
heading(doc,'8.4 Evaluation Protocol',2)
bullets(doc,['Prepare CCTV like test videos with consent or appropriate authorization.','Annotate person boxes, track IDs, visible attributes, action intervals and query relevance.','Split data by video or camera location so the same scene is not present in both training and test data.','Run the full pipeline with fixed configuration and record hardware, model versions, frame sample rate and thresholds.','Compare predicted records to annotations, calculate metrics, review false positives and false negatives, and report limitations.','Evaluate difficult conditions separately: low light, blur, occlusion, crowds, small persons and motion.'])

heading(doc,'9 Current Results and Limitations',1)
heading(doc,'9.1 Implemented and Testable Components',2)
para(doc,'The application has been built as a working vertical slice. The upload endpoint stores valid video files and reads their actual metadata. The background pipeline reports state through processing_jobs. YOLOv8 plus ByteTrack is invoked for person tracking. Frames, detections, tracks, and evidence images are stored in SQLite and storage directories. The UI presents the original footage, processing progress, a deterministic query form, temporal result cards and evidence seeking.')
heading(doc,'9.2 Model Dependent Limitations',2)
table(doc,['Component','Current Constraint','Correct Handling'],[
['Qwen3 VL','Requires compatible Transformers runtime, model availability, CUDA enabled PyTorch and adequate GPU or RAM','Show an explicit warning; do not save invented attributes'],
['VideoMAE','Requires a compatible fine tuned checkpoint and exact label mapping','Skip actions; do not return action evidence'],
['Local laptop deployment','RTX 3050 Laptop GPU has 4 GB VRAM and PyTorch CUDA was not available in the observed environment','Use CUDA enabled PyTorch for YOLO; use larger local or cloud GPU for dependable Qwen'],
['CCTV quality','Blur, low light, occlusion, scale and compression affect all vision modules','Use confidence, unknown values, original video review and labelled evaluation']],[1.45,2.6,2.15])
heading(doc,'9.3 Ethical and Operational Limits',2)
para(doc,'The system is an evidence discovery aid, not a decision maker. Results may contain detection, tracking, attribute or action errors. Access controls, upload validation, safe filenames, user ownership checks and clear warnings are required. Investigators should use the original video, timestamps and surrounding context to validate any candidate result.')

heading(doc,'10 Conclusion and Future Work',1)
para(doc,'AI Detective demonstrates a practical architecture for explainable CCTV investigation. Its central contribution is the connection between visual processing and reviewable evidence: a query maps to structured conditions, matching records retain their timestamps and track IDs, and the user is taken back to the exact original video interval. The design deliberately distinguishes real pipeline output from unavailable model dependent output.')
heading(doc,'10.1 Future Enhancements',2)
bullets(doc,['Install CUDA enabled PyTorch and validate GPU inference for YOLO and Qwen3 VL.','Select the best quality crop per track using detection confidence, crop size and sharpness before Qwen inference.','Cache validated Qwen outputs and action windows to avoid repeated model execution.','Fine tune VideoMAE Large with authorised CCTV action clips and a documented label mapping.','Add action window inference, temporal smoothing, threshold calibration and action interval persistence.','Create an annotated evaluation dataset and publish measured detection, tracking, attribute, action and retrieval metrics.','Add report export, investigation cases, privacy audit logging and multi camera time synchronization.'])

heading(doc,'References',1)
refs=[
'[1] Ultralytics. YOLO Object Detection and Multi Object Tracking Documentation. https://docs.ultralytics.com/.',
'[2] Zhang, Y., Sun, P., Jiang, Y., et al. ByteTrack: Multi Object Tracking by Associating Every Detection Box. ECCV, 2022.',
'[3] Qwen Team. Qwen3 VL 2B Instruct Model Card. Hugging Face, 2025. https://huggingface.co/Qwen/Qwen3-VL-2B-Instruct.',
'[4] Tong, Z., Song, Y., Wang, J., and Wang, L. VideoMAE: Masked Autoencoders Are Data Efficient Learners for Self Supervised Video Pre Training. NeurIPS, 2022.',
'[5] Hugging Face. Transformers Documentation for VideoMAE. https://huggingface.co/docs/transformers/model_doc/videomae.',
'[6] FastAPI. FastAPI Documentation. https://fastapi.tiangolo.com/.',
'[7] OpenCV. Open Source Computer Vision Library Documentation. https://docs.opencv.org/.',
'[8] SQLAlchemy. SQLAlchemy Documentation. https://docs.sqlalchemy.org/.']
for x in refs: para(doc,x)

doc.add_page_break(); heading(doc,'Appendix A API Summary',1)
table(doc,['Method','Endpoint','Purpose'],[
['POST','/api/videos/upload','Upload and validate a video'],
['GET','/api/videos','List locally available videos'],
['GET','/api/videos/{video_id}/file','Stream original video'],
['POST','/api/videos/{video_id}/process','Create a background processing job'],
['GET','/api/jobs/{job_id}','Read actual processing stage and percentage'],
['POST','/api/search','Parse query and retrieve matching temporal evidence'],
['GET','/api/frames/{frame_id}/evidence','Return real evidence frame with bounding box'],
['POST','/api/results/{video_id}/{track_id}/generate-video','Generate overlay evidence video']],[.8,2.3,3.3])
heading(doc,'Appendix B Environment Configuration',1)
table(doc,['Variable','Purpose','Example'],[
['YOLO_MODEL','YOLO model selection','yolov8m.pt'],
['FRAME_SAMPLE_FPS','Approximate frames processed per second','3'],
['QWEN_MODEL','Qwen3 VL model identifier or local path','Qwen/Qwen3-VL-2B-Instruct'],
['ACTION_MODEL_PATH','Compatible VideoMAE action checkpoint path','C:\\models\\videomae_actions'],
['ACTION_WINDOW_FRAMES','Frames in one action window','16'],
['ACTION_STRIDE','Frames advanced between action windows','8'],
['DEVICE','Runtime selection','auto or cuda'],
['DATABASE_URL','SQLite connection string','sqlite:///./ai_detective.db']],[1.5,2.6,2.3])

doc.core_properties.title='AI Detective Project Report'
doc.core_properties.subject='Vision Language CCTV Investigation and Search System'
doc.core_properties.author='AI Detective Project Team'
doc.save(OUT)
print(OUT)
